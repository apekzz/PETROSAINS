"""Extract the programme catalogue and train ProgrammeRanker.

Run from inventory_app:
  python backend/train_consultant.py
  python backend/train_consultant.py --check
"""

from __future__ import annotations

import json
import os
import sys

import numpy as np
import torch

import consultant_model as model

XLSX = os.path.expanduser("~/Downloads/DATASET_PROGRAMME CATALOGUE.xlsx")
CHECK_CLIMATE = "A youth club wants a programme on climate and sustainable technology for secondary students."
CHECK_SHORT = "A thirty minute hands-on slot."
CHECK_YOUNG = "Early years session on creative science for young learners with colour mixing."
DEMO_TEXT = "A school requests a sustainability programme for 200 students."


def _snake(name: str) -> str:
    import re

    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def _rows(ws):
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None and str(cell).strip() for cell in row):
            yield row


def _sheet_records(ws) -> list[dict]:
    records = []
    header = None
    for row in _rows(ws):
        cells = ["" if cell is None else str(cell).strip() for cell in row]
        if header is None:
            header = [_snake(cell) for cell in cells]
            continue
        item = {}
        for key, value in zip(header, row):
            if not key:
                continue
            item[key] = value if not isinstance(value, str) else value.strip()
        if any(item.values()):
            records.append(item)
    return records


def _kits(ws) -> list[dict]:
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    header_at = None
    item_col = None
    pack_col = None
    for index, row in enumerate(rows[:20]):
        cells = ["" if cell is None else str(cell).strip().lower() for cell in row]
        if not any(cell == "item" or cell.startswith("item") for cell in cells):
            continue
        header_at = index
        for col, cell in enumerate(cells):
            if cell == "item" or cell.startswith("item"):
                item_col = col
            if "ready pack" in cell or cell == "number of ready pack needed":
                pack_col = col
        break
    if header_at is None or item_col is None:
        return []
    lines = []
    for row in rows[header_at + 1 :]:
        if item_col >= len(row) or row[item_col] is None:
            continue
        name = str(row[item_col]).strip().split("\n")[0][:180]
        if not name or name.lower() in {"item", "nil"}:
            continue
        packs = None
        if pack_col is not None and pack_col < len(row) and row[pack_col] is not None:
            packs = model._num(row[pack_col])
        lines.append({"name": name, "packs": packs})
    return lines


def extract() -> dict:
    import openpyxl

    book = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    offerings = [row for row in _sheet_records(book["Offerings_Master"]) if row.get("offering_id")]
    mappings = [row for row in _sheet_records(book["Theme_Objective_Mapping"]) if row.get("mapping_id")]
    rules = [row for row in _sheet_records(book["Constraint_Rules"]) if row.get("rule_id")]
    kits = {}
    for offering in offerings:
        offering_id = str(offering["offering_id"])
        sheet_name = next((name for name in book.sheetnames if name.startswith(offering_id + "_") or name.startswith(offering_id + " ")), None)
        kits[offering_id] = _kits(book[sheet_name]) if sheet_name else []
    book.close()
    payload = {"offerings": offerings, "mappings": mappings, "rules": rules, "kits": kits}
    os.makedirs(os.path.dirname(model.CATALOGUE_PATH), exist_ok=True)
    with open(model.CATALOGUE_PATH, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2, default=str)
    if not os.path.isfile(model.BOOKINGS_PATH):
        with open(model.BOOKINGS_PATH, "w", encoding="utf-8") as handle:
            json.dump({"booked_offering_ids": ["ACT-018"]}, handle, indent=2)
    return payload


def _base_request(text: str, offering: dict | None = None) -> dict:
    age = None
    count = None
    duration = None
    venue = "indoor"
    electricity = "unknown"
    internet = "unknown"
    water = "unknown"
    if offering is not None:
        low, _high = model.age_bounds(offering.get("recommended_age"))
        age = int(low) if low is not None else None
        count = int(model._num(offering.get("min_participants")) or 12)
        delivery = model._num(offering.get("standard_duration_min"))
        setup = model._num(offering.get("setup_time_min")) or 0
        if delivery is not None:
            duration = int(delivery + setup)
        place = str(offering.get("indoor_outdoor") or "").lower()
        venue = "outdoor" if "outdoor" in place and "indoor" not in place else "indoor"
        electricity = model._flag(offering.get("electricity_required"))
        if electricity == "optional":
            electricity = "no"
        internet = model._flag(offering.get("internet_required"))
        if internet == "yes":
            internet = "yes"
        water = "yes" if model._flag(offering.get("water_required")) == "yes" else "unknown"
    return {
        "text": text,
        "age": age,
        "count": count if count != 200 else 24,
        "duration": duration,
        "venue": venue,
        "electricity": electricity,
        "internet": internet,
        "water": water,
        "budget": "unknown",
        "accessibility": "",
        "positive_ids": [],
    }


def build_samples(catalogue: dict) -> list[dict]:
    samples = []
    by_id = {row["offering_id"]: row for row in catalogue["offerings"]}
    for offering in catalogue["offerings"]:
        text = (
            f"Programme about {offering.get('suitable_themes') or offering.get('activity_title')}. "
            f"Audience {offering.get('audience_types') or ''}. {offering.get('suitable_objectives') or ''}"
        )
        sample = _base_request(text, offering)
        sample["positive_ids"] = [offering["offering_id"]]
        samples.append(sample)
    for mapping in catalogue["mappings"]:
        ids = [mapping.get("primary_offering_id"), mapping.get("secondary_offering_id")]
        ids = [item for item in ids if item in by_id]
        if not ids:
            continue
        text = (
            f"Theme {mapping.get('theme')}. Objective {mapping.get('stakeholder_objective')}. "
            f"Audience {mapping.get('suitable_audience')}. Concepts {mapping.get('related_concepts')}."
        )
        sample = _base_request(text, by_id[ids[0]])
        sample["text"] = text
        sample["positive_ids"] = ids
        samples.append(sample)
    school = _base_request(
        "A school requests a sustainability programme for a single class.",
        by_id.get("ACT-018"),
    )
    school["text"] = "A school requests a sustainability programme for a single class."
    school["count"] = 24
    school["positive_ids"] = ["ACT-018", "ACT-003", "ACT-019"]
    samples.append(school)
    for offering_id in ("ACT-009", "ACT-018", "ACT-001"):
        offering = by_id.get(offering_id)
        if offering is None:
            continue
        short = _base_request(
            f"Half-hour morning window for {offering.get('activity_title')}.",
            offering,
        )
        short["duration"] = 30
        short["allow_empty"] = True
        short["positive_ids"] = [offering_id]
        samples.append(short)
    kept = []
    for sample in samples:
        if sample.get("count") == 200 and "sustainability" in sample["text"].lower():
            continue
        if sample["text"] in {CHECK_CLIMATE, CHECK_YOUNG, DEMO_TEXT, CHECK_SHORT}:
            continue
        legal = [item for item in sample["positive_ids"] if item in by_id and not model.blocking_rule(sample, by_id[item])]
        if not legal and not sample.get("allow_empty"):
            continue
        sample = dict(sample)
        sample["positive_ids"] = legal
        kept.append(sample)
    return kept


def train() -> None:
    catalogue = model.load_catalogue()
    offerings = catalogue["offerings"]
    samples = build_samples(catalogue)
    documents = [model.offering_document(row) for row in offerings]
    offering_vecs = model.encode_texts(documents)
    request_vecs = model.encode_texts([model.request_document(sample) for sample in samples])
    features = []
    labels = []
    for sample_index, sample in enumerate(samples):
        wanted = set(sample["positive_ids"])
        for offering_index, offering in enumerate(offerings):
            features.append(
                model.pair_features(sample, offering, request_vecs[sample_index], offering_vecs[offering_index], catalogue)
            )
            blocked = model.blocking_rule(sample, offering) is not None
            labels.append(1.0 if offering["offering_id"] in wanted and not blocked else 0.0)
    features_t = torch.from_numpy(np.stack(features))
    labels_t = torch.tensor(labels, dtype=torch.float32)
    ranker = model.ProgrammeRanker(features_t.shape[1])
    loss_fn = torch.nn.BCEWithLogitsLoss(pos_weight=torch.tensor([6.0]))
    opt = torch.optim.Adam(ranker.parameters(), lr=1e-3)
    ranker.train()
    for epoch in range(150):
        opt.zero_grad()
        loss = loss_fn(ranker(features_t), labels_t)
        loss.backward()
        opt.step()
        if epoch in {0, 149} or epoch % 30 == 0:
            print(f"epoch {epoch} loss {loss.item():.4f} samples {len(samples)}")
    os.makedirs(os.path.dirname(model.WEIGHTS_PATH), exist_ok=True)
    torch.save(
        {
            "state_dict": ranker.state_dict(),
            "in_dim": int(features_t.shape[1]),
            "offering_ids": [row["offering_id"] for row in offerings],
            "offering_embeddings": offering_vecs,
            "encoder_name": model.ENCODER_NAME,
        },
        model.WEIGHTS_PATH,
    )
    model._PACK = None
    print("saved", model.WEIGHTS_PATH)


def _clear(payload: dict) -> dict:
    return {
        "text": payload["text"],
        "age": payload.get("age"),
        "count": payload.get("count"),
        "duration": payload.get("duration"),
        "venue": payload.get("venue") or "unknown",
        "electricity": payload.get("electricity") or "unknown",
        "internet": payload.get("internet") or "unknown",
        "water": payload.get("water") or "unknown",
        "budget": payload.get("budget") or "unknown",
        "accessibility": payload.get("accessibility") or "",
    }


def run_checks() -> None:
    catalogue_ids = {row["offering_id"] for row in model.load_catalogue()["offerings"]}
    climate = _clear(
        {
            "text": CHECK_CLIMATE,
            "age": 16,
            "count": 20,
            "duration": 360,
            "venue": "indoor",
            "electricity": "yes",
            "internet": "unknown",
            "water": "unknown",
            "budget": "unknown",
        }
    )
    ranked = {row["offering_id"]: row for row in model.rank(climate)}
    green = max(ranked[item]["score"] or -1 for item in ("ACT-018", "ACT-019", "ACT-003"))
    assert green > (ranked["ACT-011"]["score"] or -1), (green, ranked["ACT-011"])
    young = _clear(
        {
            "text": CHECK_YOUNG,
            "age": 5,
            "count": 12,
            "duration": 120,
            "venue": "indoor",
            "electricity": "no",
            "internet": "no",
            "water": "yes",
            "budget": "low",
        }
    )
    young_rank = model.rank(young)
    top = next(row for row in young_rank if not row["blocked"])
    assert top["offering_id"] != "ACT-018"
    assert young_rank[[row["offering_id"] for row in young_rank].index("ACT-018")]["blocked"] == "RULE-001"
    demo = model.advise(
        _clear(
            {
                "text": DEMO_TEXT,
                "age": None,
                "count": 200,
                "duration": None,
                "venue": "unknown",
                "electricity": "unknown",
                "internet": "unknown",
                "water": "unknown",
                "budget": "unknown",
            }
        ),
        [],
    )
    ids = [row["offering_id"] for row in demo["recommendations"]]
    assert ids and set(ids) <= catalogue_ids
    assert "ACT-018" not in ids
    assert set(ids) & {"ACT-003", "ACT-019"}, ids
    assert demo["alternative"] and demo["alternative"]["offering_id"] == "ACT-018"
    assert demo["missing_questions"]
    assert demo["rotations"]
    assert all(item.startswith("Proposed Enhancement:") for item in demo["enhancements"])
    short_case = model.advise(
        _clear(
            {
                "text": CHECK_SHORT,
                "age": 12,
                "count": 15,
                "duration": 30,
                "venue": "indoor",
                "electricity": "yes",
                "internet": "unknown",
                "water": "unknown",
                "budget": "unknown",
            }
        ),
        [],
    )
    catalogue = {row["offering_id"]: row for row in model.load_catalogue()["offerings"]}
    for row in short_case["recommendations"]:
        offering = catalogue[row["offering_id"]]
        setup = model._num(offering.get("setup_time_min")) or 0
        delivery = model._num(offering.get("standard_duration_min")) or 0
        assert setup + delivery <= 30, row
    print("checks ok", ids, "alt", demo["alternative"]["offering_id"], "short", [row["offering_id"] for row in short_case["recommendations"]])


def main() -> None:
    if "--check" not in sys.argv:
        if os.path.isfile(XLSX):
            extract()
        elif not os.path.isfile(model.CATALOGUE_PATH):
            raise SystemExit(f"missing catalogue source {XLSX}")
        train()
    run_checks()


if __name__ == "__main__":
    main()
